#!/usr/bin/env python3
"""ЛР-1 по ТАУ (Великанов/Мартынова, 2019): «Экспериментальное определение частотных характеристик
динамического звена» — комплект отчёта для конкретного студента.

Звено: W(p) = K / (a2 p^2 + a1 p + 1), вход u = Um cos(wt). В LINSYS студент для каждого периода Ti снимает
амплитуду выхода Xm и сдвиг ΔT; дальше считает w, A, φ, L и строит АЧХ, ФЧХ, АФЧХ, ЛАЧХ+ЛФЧХ.

Скрипт делает для каждого человека свой набор чисел:
  * вариант (a2, a1, K, Um) — от преподавателя; если не задан — случайный правдоподобный;
  * «почерк» студента — свои коэффициенты из рекомендованных диапазонов методички
    (T0 = 80…100·√a2, T1 = 12…13·√a2, Ti+1 = 0.80…0.85·Ti, 13–15 измерений, округление √a2);
  * «измерения курсором» — теория + шум ≈0.3 %, все величины с 3 знаками после запятой (--dec 3, по умолчанию;
    --dec 0 — прежнее округление «как с экрана»: сотые/тысячные/3 значащие).
Всё детерминировано: одинаковое ФИО (или --seed) → одинаковые числа.

Команды:
  lab1.py make  --name "Иванов И.И." [--group АТ-23-01] [--a2 2.9 --a1 0.85 --K 2 --Um 0.8] --out DIR
                [--seed N] [--dec 3] [--lach-variants] [--no-mathcad] [--teacher "…"] [--title-page]
  lab1.py batch people.csv --out DIR [--no-mathcad] [--lach-variants]
        CSV (UTF-8, ; или ,): name;group;a2;a1;K;Um[;seed][;dec] — пустые a2…Um → случайный вариант
  lab1.py check DIR/<файл>.json      — пересчитать таблицу и сверить с теорией (контроль)

Результат в DIR/<Фамилия>/: <Фамилия>_ЛР1.xmcd (пересчитан в Mathcad 13), <Фамилия>_ЛР1_графики.docx
(4 рисунка, альбомный А4), <Фамилия>_ЛР1_тетрадь.docx (данные, цепочка периодов, расчёты, таблица 1,
аналитические выражения); служебное/: data.json, plots/*.png, контрольные_вопросы.docx.
С --lach-variants ещё: _графики_разные_нули.docx, _графики_общий_ноль_дробный_шаг.docx, _ЛАФЧХ_6_вариантов.docx.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import random
import re
import shutil
import struct
import subprocess
import sys
import zlib
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL / "assets" / "lab1_template.xmcd"

GRID_RGB = (0, 255, 0)
TITLE = "Экспериментальное определение частотных характеристик динамического звена"
FIGS = [
    ("ach", "Амплитудная частотная характеристика (АЧХ)"),
    ("fch", "Фазовая частотная характеристика (ФЧХ)"),
    ("afch", "Амплитудно-фазовая частотная характеристика (АФЧХ)"),
    ("lach", "Логарифмическая амплитудная и логарифмическая фазовая частотные характеристики (ЛАЧХ и ЛФЧХ)"),
]


# ============================================================================ числа
def fnum(x, nd=3):
    """Число для тетради: до nd знаков после запятой, без хвостовых нулей, десятичная запятая."""
    s = f"{x:.{nd}f}".rstrip("0").rstrip(".")
    if s in ("-0", ""):
        s = "0"
    return s.replace(".", ",")


def mnum(x, sig=10):
    """Число для Mathcad (точка, без экспоненты)."""
    s = f"{x:.{sig}g}"
    if "e" in s:
        s = f"{x:.12f}".rstrip("0").rstrip(".")
    return s


def seed_of(name):
    return zlib.crc32(name.strip().lower().encode("utf-8"))


def meas_round(x, kind="Xm"):
    """Отсчёт курсором как в тетрадях: ≥ 1 — сотые (1.61, 2.45), 0.1…1 — тысячные (0.812),
    < 0.1 — три значащие (0.0261), иначе в хвосте АЧХ ошибка округления доходит до 5 %."""
    return round(x, 2) if x >= 1 else round(x, 3) if x >= 0.1 else float(f"{x:.3g}")


def freq_resp(K, a2, a1, w):
    W = K / complex(1 - a2 * w * w, a1 * w)
    ph = math.degrees(math.atan2(W.imag, W.real))  # для звена 2-го порядка лежит в (-180, 0]
    if ph > 0:
        ph -= 360
    return W, abs(W), ph


def random_variant(rng):
    while True:
        a2 = round(rng.uniform(2.0, 4.0), 1)
        a1 = round(rng.uniform(0.6, 1.2) / 0.05) * 0.05
        xi = a1 / (2 * math.sqrt(a2))
        if 0.15 <= xi <= 0.33:  # ярко выраженный резонанс, как в реальных вариантах
            break
    K = round(rng.uniform(1.0, 3.0), 1)
    Um = round(rng.uniform(0.5, 1.6), 1)
    return dict(a2=a2, a1=round(a1, 2), K=K, Um=Um)


def make_data(name, group="", a2=None, a1=None, K=None, Um=None, seed=None, dec=None, **extra):
    """dec=3 — все величины (τ, T, Xm, ΔT, ω, A, φ, L) ровно с 3 знаками после запятой, производные считаются из
    уже округлённых (в тетради «20 lg 0,011 = −39,172» сходится при проверке на калькуляторе)."""
    seed = seed_of(name) if seed in (None, "") else int(seed)
    dec = int(dec) if dec not in (None, "") else None
    rng = random.Random(seed)
    given = None not in (a2, a1, K, Um)
    var = dict(a2=float(a2), a1=float(a1), K=float(K), Um=float(Um)) if given else random_variant(rng)
    if given:  # «почерк» не должен зависеть от того, был ли вариант случайным
        random_variant(rng)
    style = dict(
        tau_dec=rng.choice([1, 2, 2]),
        c0=rng.choice([80, 85, 90, 90, 95, 100]),
        c1=rng.choice([12, 12.5, 12.5, 13]),
        q=rng.choice([0.80, 0.81, 0.82, 0.83, 0.84, 0.85]),
        n=rng.choice([13, 14, 15, 15]),
    )
    a2, a1, K, Um = var["a2"], var["a1"], var["K"], var["Um"]
    if dec:
        tau = round(math.sqrt(a2), dec)
        T = [round(style["c0"] * tau, dec), round(style["c1"] * tau, dec)]
        while len(T) < style["n"]:
            T.append(round(style["q"] * T[-1], dec))
    else:
        tau = round(math.sqrt(a2), style["tau_dec"])
        T = [round(style["c0"] * tau, 1), round(style["c1"] * tau, 2)]
        while len(T) < style["n"]:
            T.append(round(style["q"] * T[-1], 2))

    rows = []
    prev = None
    for Ti in T:
        w = 2 * math.pi / Ti
        _, A_t, ph_t = freq_resp(K, a2, a1, w)
        Xm_t, dT_t = Um * A_t, -ph_t / 360 * Ti
        for _attempt in range(200):
            # ошибка отсчёта курсором ≈0.3–0.4 % (не больше 1.5 %); точки идут монотонно, как у теории:
            # фаза только убывает, A растёт до резонанса и потом падает — иначе заметен «скачок» на графике
            ex, et = rng.gauss(0, 0.003), rng.gauss(0, 0.003)
            if abs(ex) >= 0.015 or abs(et) >= 0.015:
                continue
            if dec:
                Xm, dT = round(Xm_t * (1 + ex), dec), round(dT_t * (1 + et), dec)
                if Xm <= 0:
                    continue
                A_i, ph_i = round(Xm / Um, dec), round(-dT / Ti * 360, dec)
            else:
                Xm, dT = meas_round(Xm_t * (1 + ex), "Xm"), meas_round(dT_t * (1 + et), "dT")
                A_i, ph_i = Xm / Um, -dT / Ti * 360
            if prev is None or (ph_i < prev[1] and (A_i - prev[0]) * (A_t - prev[2]) > 0):
                break
        w_i = 2 * math.pi / Ti
        L_i = 20 * math.log10(round(A_i, dec) if dec else A_i)
        prev = (A_i, ph_i, A_t)
        rows.append(dict(T=Ti, Xm=Xm, dT=dT, w=round(w_i, 3), A=round(A_i, 3), phi=round(ph_i, 3),
                         L=round(L_i, 3), A_theory=A_t, phi_theory=ph_t))

    xi = a1 / (2 * math.sqrt(a2))
    wn = 1 / math.sqrt(a2)
    an = dict(T=math.sqrt(a2), xi=xi, wn=wn)
    if xi < 1 / math.sqrt(2):
        an["wr"] = wn * math.sqrt(1 - 2 * xi * xi)
        an["Amax"] = K / (2 * xi * math.sqrt(1 - xi * xi))
        an["M"] = an["Amax"] / K
    exp_max = max(rows, key=lambda r: r["A"])
    return dict(name=name.strip(), group=group, seed=seed, dec=dec, variant=var, style=style, tau=tau,
                rows=rows, analytic=an, exp_peak=dict(w=exp_max["w"], A=exp_max["A"]), **extra)


# ============================================================================ Mathcad
def find_mathcad_scripts():
    cands = [os.environ.get("MATHCAD13_SKILL"), Path.home() / ".claude" / "skills" / "mathcad13",
             SKILL.parent / "mathcad13", SKILL.parent.parent / "Mathcad 13" / "mathcad13",
             # репозиторий claude_skills: Отчёты/ТАУ лабы/tau-labs → Программы/Mathcad 13/mathcad13
             SKILL.parent.parent.parent / "Программы" / "Mathcad 13" / "mathcad13"]
    for c in cands:
        if c and (Path(c) / "scripts" / "xmcd.py").exists():
            return Path(c) / "scripts"
    raise FileNotFoundError("не найден скилл mathcad13 (scripts/xmcd.py); задайте MATHCAD13_SKILL")


def nice_step(span, max_div):
    for s in (0.05, 0.1, 0.2, 0.25, 0.5, 1, 2, 2.5, 5, 10, 20, 25, 50):
        if span / s <= max_div:
            return s
    return 100


def ceil_to(x, s):
    return math.ceil(x / s - 1e-9) * s


def plot_limits(d):
    """Пределы и число делений сетки для 4 графиков (знаки как в шаблоне: минусы уже стоят в дереве)."""
    v, rows = d["variant"], d["rows"]
    K, a2, a1 = v["K"], v["a2"], v["a1"]
    wmax = rows[-1]["w"]
    ws = [wmax * i / 2000 for i in range(1, 2001)]
    resp = [freq_resp(K, a2, a1, w) for w in ws]
    A_all = [r[1] for r in resp] + [r["A"] for r in rows]
    lim = {}
    xs = nice_step(wmax * 1.05, 30)
    xtop = ceil_to(wmax * 1.02, xs)
    ys = nice_step(max(A_all) * 1.05, 50)
    ytop = ceil_to(max(A_all) * 1.03, ys)
    lim["ach"] = dict(nums=[ytop, 0, xtop, 0], grid=(round(xtop / xs), round(ytop / ys)))
    lim["fch"] = dict(nums=[0, 180, xtop, 0], grid=(round(xtop / xs), 36))
    # АФЧХ: одинаковый масштаб по осям (требование методички) — один шаг сетки и пропорции рамки
    re_all = [r[0].real for r in resp] + [K]
    im_all = [r[0].imag for r in resp]
    s = nice_step(max(max(re_all) - min(re_all), -min(im_all)) * 1.1, 30)
    xr, xl, yb = ceil_to(max(re_all) * 1.03, s), ceil_to(-min(re_all) * 1.05, s), ceil_to(-min(im_all) * 1.03, s)
    lim["afch"] = dict(nums=[0, yb, xr, xl], grid=(round((xr + xl) / s), round(yb / s)), aspect=(xr + xl) / yb)
    # ЛАЧХ+ЛФЧХ: 24 деления на обеих осях Y (ЛФЧХ: -180…60 по 10°), X — целые декады
    Ls = [20 * math.log10(a) for a in A_all if a > 0]
    lo_w = min(r["w"] for r in rows)
    x_lo, x_hi = 10 ** math.floor(math.log10(lo_w * 0.9)), 10 ** math.ceil(math.log10(wmax * 1.1))
    Lmin = min(r["L"] for r in rows)  # кривая правее последней точки может уйти под рамку — как в образце
    for st in (1, 2, 2.5, 4, 5, 10):
        top = ceil_to(max(Ls) + st / 2, st)
        if top - 24 * st <= Lmin - st / 2:
            break
    lim["lach"] = dict(nums=[top, 24 * st - top, 60, 180, x_hi, x_lo], grid=None, Lstep=st)
    for k in lim:
        lim[k]["nums"] = [float(f"{x:.6g}") for x in lim[k]["nums"]]
    return lim


def patch_plot(blob, nums, grid=None, aspect=None):
    """Заменить пользовательские пределы осей (листья-числа без бита «авто» 0x02) по порядку дерева,
    число делений сетки и (для АФЧХ) пропорции рамки."""
    i = blob.find(b"tree")
    head, tail = blob[:i], blob[i:]
    leaf = re.compile(rb"(\x40\x40\x02\x0f\x00\x00)([\x00-\xff])(\x40[\x00-\xff]\x00\x00|[^\x40]\x00\x00)"
                      rb"([\x02-\x40])\x01([\x01-\x3f])([-0-9.eE]+)\x00")
    out, pos, k = bytearray(), 0, 0
    for m in leaf.finditer(tail):
        side = m.group(2)[0]
        if side & 0x02 or (side & 0x30) != 0x30:  # авто-значение или не число
            continue
        if k >= len(nums):
            break
        v = nums[k]
        txt = (f"{v:.6g}" if abs(v) >= 1e-4 else "0").encode()
        out += tail[pos:m.start()] + m.group(1) + m.group(2) + m.group(3) + bytes([len(txt) + 1, 1, len(txt)]) + txt + b"\x00"
        pos = m.end()
        k += 1
    if k != len(nums):
        raise RuntimeError(f"в графике {k} пользовательских пределов, ожидалось {len(nums)}")
    data = bytearray(head + bytes(out) + tail[pos:])
    p = data.find(b"axisFormat") + len(b"axisFormat") + 10
    for off in (p, p + 22):  # зелёная сетка на всех графиках (axisFormat [14..16] = RGB)
        data[off + 14:off + 17] = bytes(GRID_RGB)
    if grid:
        gx, gy = grid
        data[p + 9] = max(2, min(99, gx))
        data[p + 22 + 9] = max(2, min(99, gy))
    if aspect:
        j = data.find(b"graphData") + len(b"graphData")
        w, h = struct.unpack("<ii", data[j + 27:j + 35])
        h = 89
        w = int(round(h * aspect))
        if w > 150:
            w, h = 150, int(round(150 / aspect))
        data[j + 27:j + 35] = struct.pack("<ii", w, h)
    return bytes(data)


def build_xmcd(d, out_path, recalc=True):
    scripts = find_mathcad_scripts()
    sys.path.insert(0, str(scripts))
    import xmcd  # noqa: E402
    import mcplot  # noqa: E402
    import xml.etree.ElementTree as ET

    text = xmcd.strip_digest(TEMPLATE.read_text(encoding="utf-8"))
    root = ET.fromstring(text.encode("utf-8"))
    v, rows = d["variant"], d["rows"]
    K, a2, a1 = v["K"], v["a2"], v["a1"]
    vec = lambda key: "[" + "; ".join(mnum(r[key]) for r in rows) + "]"
    D = f"(((-{mnum(a2)})*w^2 + 1)^2 + {mnum(a1 * a1)}*w^2)"
    assigns = {
        "Um": mnum(v["Um"]), "a2": mnum(a2), "a1": mnum(a1), "K": mnum(K),
        "T": vec("T"), "Xm": vec("Xm"), "ΔT": vec("dT"),
        "U": f"((-{mnum(K * a2)})*w^2 + {mnum(K)})/{D}",
        "V": f"{mnum(K * a1)}*(w/{D})",
    }
    ML = xmcd.ML
    done = set()
    for dfn in root.iter(f"{{{ML}}}define"):
        lhs = list(dfn)[0]
        lname = xmcd.ids_text(lhs) if xmcd.local(lhs.tag) == "id" else (
            xmcd.ids_text(list(lhs)[0]) if xmcd.local(lhs.tag) == "function" else None)
        if lname in assigns and lname not in done:
            node = xmcd.Parser(assigns[lname]).range_or_expr()
            new_el = ET.fromstring(f'<root xmlns:ml="{ML}">' + xmcd.to_xml(node) + "</root>")[0]
            dfn.remove(list(dfn)[1])
            dfn.insert(1, new_el)
            done.add(lname)
    miss = set(assigns) - done
    if miss:
        raise RuntimeError(f"в шаблоне не найдены определения: {miss}")

    WS = "{" + xmcd.WS + "}"
    regions = root.find(WS + "regions")
    for reg in list(regions):  # L2(w) := 20·log(φ1(w)) в образце — бессмысленная строка, убираем
        if re.search(r"L2", "".join(reg.itertext())) and reg.find(WS + "math") is not None:
            ids = [xmcd.ids_text(e) for e in reg.iter(f"{{{ML}}}id")]
            if ids and ids[0] == "L2":
                regions.remove(reg)
    lim = plot_limits(d)
    items = xmcd.binary_items(root)
    plot_regs = [r for r in regions if r.find(WS + "plot") is not None]
    if len(plot_regs) != 4:
        raise RuntimeError("в шаблоне ожидалось 4 графика")
    for (key, _), reg in zip(FIGS, plot_regs):
        it = items[reg.find(WS + "plot").get("item-idref")]
        L = lim[key]
        blob = patch_plot(xmcd.item_bytes(it), L["nums"], L.get("grid"), L.get("aspect"))
        it.text = mcplot.encode(blob)
        it.set("content-encoding", "gzip")
        reg.set("tag", "plot_" + key)
        if L.get("aspect"):
            j = blob.find(b"graphData") + len(b"graphData")
            w, h = struct.unpack("<ii", blob[j + 27:j + 35])
            reg.set("width", "%.1f" % (w * 7.0 + 30))
            reg.set("height", "%.1f" % (h * 7.0 + 30))
    raw = out_path.with_suffix(".raw.xmcd")
    set_sheet_identity(root, WS, d)
    raw.write_text('<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
                   + ET.tostring(root, encoding="unicode"), encoding="utf-8", newline="\n")
    if not recalc:
        shutil.move(raw, out_path)
        return dict(recalc=False, limits=lim)
    import mcsheet  # noqa: E402
    r = mcsheet.mc(["-Path", str(raw), "-Regions", "-Meta", f"Author={d['name']};Title=;Description=",
                    "-SaveAs", str(out_path)])
    raw.unlink(missing_ok=True)
    if not r.get("ok"):
        raise RuntimeError("Mathcad: " + json.dumps(r, ensure_ascii=False)[:1500])
    # Mathcad при сохранении пишет revisedBy = пользователь Windows — ставим того же автора
    txt = out_path.read_text(encoding="utf-8")
    txt = re.sub(r"<revisedBy>[^<]*</revisedBy>|<revisedBy\s*/>", f"<revisedBy>{d['name']}</revisedBy>", txt)
    txt = re.sub(r"<author>[^<]*</author>|<author\s*/>", f"<author>{d['name']}</author>", txt)
    out_path.write_text(re.sub(r"<\?validation-md5-digest[^?]*\?>\s*", "", txt), encoding="utf-8", newline="\n")
    return dict(recalc=True, errors=r.get("error_count", 0), error_list=r.get("errors", []), limits=lim,
                regions=r.get("region_count"))


def set_sheet_identity(root, WS, d):
    """Шаблон — чужой лист: в метаданных e-mail автора образца и его documentID (одинаковый у всех копий —
    по нему видно копирование). Ставим автора = студент и свои ID (детерминированно от seed)."""
    import uuid
    md = root.find(WS + "metadata")
    if md is None:
        return
    ud = md.find(WS + "userData")
    if ud is not None:
        for tag in ("title", "description", "company", "keywords"):
            el = ud.find(WS + tag)
            if el is not None:
                el.text = None
        for tag in ("author", "revisedBy"):
            el = ud.find(WS + tag)
            if el is not None:
                el.text = d["name"]
    ii = md.find(WS + "identityInfo")
    if ii is not None:
        ns = uuid.UUID(int=d["seed"] * 7919 + 12345)
        for tag in ("documentID", "versionID"):
            el = ii.find(WS + tag)
            if el is not None:
                el.text = str(uuid.uuid5(ns, tag + d["name"]))
        el = ii.find(WS + "revision")
        if el is not None:
            el.text = "1"


# ============================================================================ Word
def _docx_base(landscape=False, author=""):
    import datetime as _dt
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm, Pt
    from docx.oxml.ns import qn

    doc = Document()
    cp = doc.core_properties  # по умолчанию python-docx пишет «generated by python-docx» — затираем
    cp.author, cp.last_modified_by = author, author
    cp.comments = cp.title = cp.subject = cp.keywords = cp.category = ""
    cp.revision = 1
    now = _dt.datetime.now().replace(microsecond=0)
    cp.created = cp.modified = now
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(5), Mm(5), Mm(5), Mm(5)
    else:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(30), Mm(15), Mm(20), Mm(20)
    for st in ("Normal", "Heading 1", "Heading 2", "Title"):
        s = doc.styles[st]
        s.font.name = "Times New Roman"
        s.font.size = Pt(14 if st == "Normal" else 16 if st == "Heading 1" else 14)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        if st != "Normal":
            s.font.color.rgb = None
            s.font.bold = True
    pf = doc.styles["Normal"].paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    return doc


def save_docx(doc, out):
    """Сохранить и заменить в docProps/app.xml «Microsoft Macintosh Word» (шаблон python-docx) на обычный Word."""
    import zipfile
    doc.save(out)
    tmp = Path(str(out) + ".tmp")
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "docProps/app.xml":
                data = data.replace(b"Microsoft Macintosh Word", b"Microsoft Office Word")
            zout.writestr(it, data)
    tmp.replace(out)


def title_page(doc, d, teacher):
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    def p(text="", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, before=0):
        par = doc.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_before = Pt(before)
        r = par.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
        return par

    uni = d.get("university") or "РГУ нефти и газа (НИУ) имени И.М. Губкина"
    p("Министерство науки и высшего образования Российской Федерации")
    p(uni)
    p(d.get("department") or "Кафедра автоматизации технологических процессов")
    p("ОТЧЁТ", bold=True, size=16, before=160)
    p("по лабораторной работе № 1", size=14)
    p("по дисциплине «Теория автоматического управления»")
    p(f"«{TITLE}»", bold=True, before=12)
    p(f"Выполнил: студент группы {d.get('group') or '________'}", align=WD_ALIGN_PARAGRAPH.RIGHT, before=100)
    p(d["name"], align=WD_ALIGN_PARAGRAPH.RIGHT)
    p(f"Проверил: {teacher or '________________'}", align=WD_ALIGN_PARAGRAPH.RIGHT, before=12)
    p(f"Москва, {d.get('year') or 2026} г.", before=150)


def graphs_docx(d, plots, out, teacher=None, with_title=False, items=None):
    """Графики вставляются в натуральную величину (ширина = ширине картинки в мм), чтобы клетка на распечатке
    была ровно того размера, что заложен в tauplot (масштаб «по миллиметровке»).
    items — свой список [(график, полная подпись)] вместо 4 рисунков FIGS (например, все варианты ЛАФЧХ)."""
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Mm, Pt

    if with_title:  # титул — книжный раздел, графики — альбомный
        doc = _docx_base(landscape=False, author=d["name"])
        title_page(doc, d, teacher)
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(5), Mm(5), Mm(5), Mm(5)
    else:
        doc = _docx_base(landscape=True, author=d["name"])
    if items is None:
        items = [(plots[key], f"Рисунок {n} – {cap}") for n, (key, cap) in enumerate(FIGS, 1) if plots.get(key)]
    for k, (p, caption) in enumerate(items):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = par.paragraph_format  # одинарный интервал: множитель 1,15 растягивает строку с картинкой
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0
        pf.page_break_before = k > 0
        par.add_run().add_picture(p["path"], width=Mm(p["width_mm"]))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_before, c.paragraph_format.line_spacing = Pt(2), 1.0
        c.add_run(caption)
    save_docx(doc, out)


QUESTIONS = [
    ("Как экспериментально определяются частотные характеристики объекта?",
     "На вход объекта подают гармонический сигнал u = Um·cos(ωt) и после окончания переходного процесса измеряют "
     "амплитуду Xm выходного гармонического сигнала и его сдвиг по времени ΔT относительно входного. Опыт повторяют "
     "для ряда частот (периодов Ti) от ω ≈ 0 до частот, где выход становится очень малым. Для каждой частоты "
     "A(ω) = Xm/Um, φ(ω) = ∓ΔT/T·360°."),
    ("Что такое амплитудно-фазовая характеристика?",
     "АФЧХ — годограф частотной передаточной функции W(jω) = U(ω) + jV(ω) = A(ω)·e^(jφ(ω)) на комплексной плоскости "
     "при изменении ω от 0 до +∞: геометрическое место концов векторов длиной A(ω), повёрнутых на угол φ(ω) от "
     "положительной вещественной полуоси. На неё наносят разметку частот."),
    ("Перечислите частотные характеристики и дайте их определения.",
     "АЧХ A(ω) = |W(jω)| — отношение амплитуд выхода и входа; ФЧХ φ(ω) = arg W(jω) — сдвиг фазы выхода относительно "
     "входа; АФЧХ W(jω) — годограф на комплексной плоскости; вещественная U(ω) = Re W(jω) и мнимая V(ω) = Im W(jω) "
     "характеристики; ЛАЧХ L(ω) = 20·lg A(ω) [дБ] и ЛФЧХ φ(ω) в логарифмическом масштабе частот."),
    ("Какой знак имеет фаза в выполняемой работе и почему?",
     "Отрицательный: выходной сигнал колебательного звена отстаёт от входного (инерционность), φ изменяется от 0 "
     "при ω → 0 до −180° при ω → ∞ и равен −90° на частоте ω = 1/√a2."),
    ("В каких координатах и в каком масштабе строятся логарифмические частотные характеристики?",
     "По оси абсцисс — частота ω в логарифмическом масштабе (lg ω, декада — изменение частоты в 10 раз), по оси "
     "ординат — L(ω) в децибелах и φ(ω) в градусах в линейном масштабе. Ось ординат проводят через произвольную "
     "точку, т.к. ω = 0 соответствует −∞. ЛАЧХ и ЛФЧХ обычно строят на одном графике с общей осью частот."),
    ("Как получить выражения для построения частотных характеристик, если известна передаточная функция?",
     "Подставить p = jω в W(p), выделить вещественную и мнимую части W(jω) = U(ω) + jV(ω) (домножив на сопряжённое "
     "знаменателю), затем A(ω) = √(U² + V²), φ(ω) = arg W(jω), L(ω) = 20·lg A(ω)."),
    ("Как получить аналитические выражения для фазовой частотной характеристики, если |arg W(jω)| > π/2?",
     "Формула φ = arctg(V/U) верна только при |φ| ≤ π/2 (U > 0). При U < 0 к арктангенсу добавляют ∓π: для нашего "
     "звена φ(ω) = −arctg(|V|/U) при U > 0, φ(ω) = −π − arctg(|V|/U) при U < 0 (при U = 0 φ = −π/2). Либо "
     "складывают фазы числителя и знаменателя: φ = −arg(1 − a2ω² + j·a1ω)."),
]


def notebook_docx(d, out):
    """Тетрадь: заголовок (номер лабы, тема, для кого) + исходные данные, расчёты, таблица 1, аналитические
    выражения. Все формулы — формулы Word (OMML через wordmath), числа с десятичной запятой."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    from wordmath import eq, tex_num as t

    doc = _docx_base(landscape=False, author=d["name"])
    v, rows, st, an = d["variant"], d["rows"], d["style"], d["analytic"]
    n = len(rows)
    K, a2, a1, Um = v["K"], v["a2"], v["a1"], v["Um"]
    dec = d.get("dec")

    def x(val, nd=3):  # измеренные и рассчитанные величины: до dec знаков, незначащие нули в расчётах убраны
        return t(val, dec) if dec else t(val, nd)  # (16,73; 0,04); в таблице 1 — ровно dec знаков

    def centered(text, bold=False, size=14, after=0):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(after)
        r = par.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
        return par

    def h(text):
        par = doc.add_heading(text, level=2)
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.space_after = Pt(4)

    def f(latex, text=None):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        if text:
            par.add_run(text + " ")
        eq(par, latex)
        return par

    centered("Лабораторная работа № 1", bold=True, size=16)
    centered(f"«{TITLE}»", bold=True)
    who = f"Выполнил: {d['name']}" + (f", группа {d['group']}" if d.get("group") else "")
    centered(who, after=6)

    h("Исходные данные")
    f(r"W(p)=\frac{K}{a_{2}p^{2}+a_{1}p+1}")
    f(rf"a_{{2}}={t(a2)};\ \ \  a_{{1}}={t(a1)};\ \ \  K={t(K)};\ \ \  U_{{m}}={t(Um)}")

    h("Расчёт периодов входного сигнала")
    tau = d["tau"]
    f(rf"\tau=\sqrt{{a_{{2}}}}=\sqrt{{{t(a2)}}}={x(tau)}\ \text{{с}}")
    f(rf"T_{{0}}={t(st['c0'])}\tau={t(st['c0'])}\cdot {x(tau)}={x(rows[0]['T'], 2)}\ \text{{с}}")
    f(rf"T_{{1}}={t(st['c1'])}\tau={t(st['c1'])}\cdot {x(tau)}={x(rows[1]['T'], 2)}\ \text{{с}}")
    q = t(st["q"], 2)
    for i in range(2, n):
        f(rf"T_{{{i}}}={q}T_{{{i - 1}}}={q}\cdot {x(rows[i - 1]['T'], 2)}={x(rows[i]['T'], 2)}\ \text{{с}}")

    h("Расчёт частотных характеристик")
    f(r"\omega_{i}=\frac{2\pi}{T_{i}}", "1) Значения частоты:")
    for i, r in enumerate(rows):
        f(rf"\omega_{{{i}}}=\frac{{2\pi}}{{{x(r['T'], 2)}}}={x(r['w'])}\ \text{{рад/с}}")
    f(r"A(\omega_{i})=\frac{X_{m\ i}}{U_{m}}", "2) Значения АЧХ:")
    for i, r in enumerate(rows):
        f(rf"A(\omega_{{{i}}})=\frac{{{x(r['Xm'], 4)}}}{{{t(Um)}}}={x(r['A'])}")
    f(r"\varphi(\omega_{i})=-\frac{\Delta T_{i}}{T_{i}}\cdot 360^{\circ}", "3) Значения ФЧХ:")
    for i, r in enumerate(rows):
        f(rf"\varphi(\omega_{{{i}}})=-\frac{{{x(r['dT'])}}}{{{x(r['T'], 2)}}}\cdot 360^{{\circ}}={{{x(r['phi'])}}}^{{\circ}}")
    f(r"L(\omega_{i})=20\ \lg\ A(\omega_{i})", "4) Значения ЛАЧХ:")
    for i, r in enumerate(rows):
        f(rf"L(\omega_{{{i}}})=20\ \lg\ {x(r['A'])}={x(r['L'])}\ \text{{дБ}}")

    h("Таблица 1")
    heads = [r"T_{i},\ \text{с}", r"X_{m\ i}", r"\Delta T_{i},\ \text{с}", r"\omega_{i},\ \text{рад/с}",
             r"A(\omega_{i})", r"\varphi(\omega_{i}),\ \text{град}", r"L(\omega_{i}),\ \text{дБ}"]
    tb = doc.add_table(rows=n + 2, cols=7)
    tb.style = "Table Grid"
    for j, hd in enumerate(heads):
        eq(tb.cell(0, j).paragraphs[0], hd)
        tb.cell(1, j).text = str(j + 1)
    for i, r in enumerate(rows):
        if dec:
            vals = [f"{r[k]:.{dec}f}".replace(".", ",") for k in ("T", "Xm", "dT", "w", "A", "phi", "L")]
        else:
            vals = [fnum(r["T"], 2), fnum(r["Xm"], 4), fnum(r["dT"]), fnum(r["w"]), fnum(r["A"]), fnum(r["phi"]),
                    fnum(r["L"])]
        for j, val in enumerate(vals):
            tb.cell(i + 2, j).text = val.replace("-", "−")
    for row in tb.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(12)

    h("Аналитические выражения")
    D = rf"\left(1-{t(a2)}\omega^{{2}}\right)^{{2}}+{t(a1 * a1, 4)}\omega^{{2}}"
    num = rf"{t(a1)}\omega"
    den = rf"1-{t(a2)}\omega^{{2}}"
    wn = t(an["wn"])
    f(rf"W(p)=\frac{{{t(K)}}}{{{t(a2)}p^{{2}}+{t(a1)}p+1}}")
    f(rf"W(j\omega)=\frac{{{t(K)}}}{{1-{t(a2)}\omega^{{2}}+j{t(a1)}\omega}}")
    f(rf"U(\omega)=\mathrm{{Re}}\ W(j\omega)=\frac{{{t(K)}\left(1-{t(a2)}\omega^{{2}}\right)}}{{{D}}}")
    f(rf"V(\omega)=\mathrm{{Im}}\ W(j\omega)=-\frac{{{t(K * a1, 4)}\omega}}{{{D}}}")
    f(rf"A(\omega)=\sqrt{{U^{{2}}(\omega)+V^{{2}}(\omega)}}=\frac{{{t(K)}}}{{\sqrt{{{D}}}}}")
    f(rf"\varphi(\omega)=-\mathrm{{arctg}}\ \frac{{{num}}}{{{den}}},\ \ \  \omega<{wn}\ \text{{рад/с}}")
    f(rf"\varphi(\omega)=-180^{{\circ}}-\mathrm{{arctg}}\ \frac{{{num}}}{{{den}}},\ \ \  \omega>{wn}\ \text{{рад/с}}")
    f(rf"L(\omega)=20\ \lg\ A(\omega)=20\ \lg\ {t(K)}-10\ \lg\left[{D}\right]")
    save_docx(doc, out)


# ============================================================================ команды
def surname(name):
    return re.split(r"[\s.]+", name.strip())[0] or "student"


LACH_VARIANTS = {  # дополнительные комплекты графиков (по просьбе): отличается только рисунок 4 — ЛАФЧХ
    "split": "разные_нули",                 # нули на разной высоте, обе шкалы с нижней линии, шаги 1|2 дБ, 5|10°
    "exact": "общий_ноль_дробный_шаг",      # нули совпадают, фаза ровно от −180°, шаг дБ 2,5 (дробный), 10°
}


def fitted(d):
    """Копия данных, где точки ωᵢ лежат ровно на расчётных кривых: A, φ, L — теоретические на тех же ωᵢ
    (округление как в таблице). Только для графиков-вариантов «с подогнанными точками»."""
    import copy
    v = d["variant"]
    nd = d.get("dec") or 3
    out = copy.deepcopy(d)
    for r in out["rows"]:
        _, A, ph = freq_resp(v["K"], v["a2"], v["a1"], r["w"])
        r["A"], r["phi"], r["L"] = round(A, nd), round(ph, nd), round(20 * math.log10(A), nd)
    return out


def lach_variant_docs(d, folder, plots, teacher=None, with_title=False):
    """Комплекты графиков с другими шкалами ЛАФЧХ рядом с основным: <Фамилия>_ЛР1_графики_<вариант>.docx,
    и <Фамилия>_ЛР1_ЛАФЧХ_6_вариантов.docx — все шкалы (nice, split, exact) с точками измерений, затем те же три
    с точками, подогнанными к расчётным кривым; у всех подпись «Рисунок 4 – …» (любой лист заменяет рисунок 4)."""
    work = Path(folder) / "служебное"
    out = []
    lach = {("nice", False): plots["lach"]}
    for mode, suffix in LACH_VARIANTS.items():
        alt = draw_plots(d, work / f"plots_{mode}", lach_mode=mode)
        lach[(mode, False)] = alt["lach"]
        pl = dict(plots, lach=alt["lach"])
        path = Path(folder) / f"{surname(d['name'])}_ЛР1_графики_{suffix}.docx"
        graphs_docx(d, pl, path, teacher, with_title)
        out.append(path)
    df = fitted(d)
    for mode in ("nice", *LACH_VARIANTS):
        lach[(mode, True)] = draw_plots(df, work / f"plots_fit_{mode}", lach_mode=mode)["lach"]
    cap = "Рисунок 4 – " + dict(FIGS)["lach"]
    order = [(m, f) for f in (False, True) for m in ("nice", *LACH_VARIANTS)]
    path = Path(folder) / f"{surname(d['name'])}_ЛР1_ЛАФЧХ_6_вариантов.docx"
    graphs_docx(d, plots, path, items=[(lach[k], cap) for k in order])
    out.append(path)
    return out


def cmd_make_one(args_d, out_root, do_mathcad=True, teacher=None, with_title=False, lach_variants=False):
    """В папке человека — только то, что сдаётся (.xmcd, графики, тетрадь); числа и картинки — в «служебное»."""
    d = make_data(**args_d)
    folder = Path(out_root) / surname(d["name"])
    folder.mkdir(parents=True, exist_ok=True)
    work = folder / "служебное"
    work.mkdir(exist_ok=True)
    for old in ("data.json", "plots"):  # раскладка прежних версий скрипта
        pth = folder / old
        if pth.is_dir():
            shutil.rmtree(pth, ignore_errors=True)
        elif pth.exists():
            pth.unlink()
    base = f"{surname(d['name'])}_ЛР1"
    report = build_xmcd(d, folder / f"{base}.xmcd") if do_mathcad else {}
    plots = draw_plots(d, work / "plots")
    (work / "data.json").write_text(json.dumps(dict(d, mathcad=report, plots=plots), ensure_ascii=False, indent=1),
                                    encoding="utf-8")
    graphs_docx(d, plots, folder / f"{base}_графики.docx", teacher, with_title)
    notebook_docx(d, folder / f"{base}_тетрадь.docx")
    questions_docx(d, work / "контрольные_вопросы.docx")
    if lach_variants:
        lach_variant_docs(d, folder, plots, teacher, with_title)
    return d, folder, dict(report, plots=plots)


def questions_docx(d, out):
    """Ответы на контрольные вопросы — для подготовки к защите (в тетрадь не переписываются)."""
    doc = _docx_base(landscape=False, author=d["name"])
    for i, (q, a) in enumerate(QUESTIONS, 1):
        doc.add_paragraph().add_run(f"{i}. {q}").bold = True
        doc.add_paragraph(a)
    save_docx(doc, out)


def draw_plots(d, outdir, lach_mode="nice"):
    """4 графика с нуля (tauplot), формат Crossed: оси со стрелками и подписями у концов, числа вдоль осей,
    АФЧХ — квадратная клетка и одинаковый шаг, разметка частот без пересечений. Вставлять в Word в натуральную
    величину. lach_mode — шкалы ЛАЧХ+ЛФЧХ:
      "nice"  — нули на одной линии, шаги 1|2 дБ и 5|10° на клетку (фаза может уйти ниже −180°) — по умолчанию;
      "exact" — нули на одной линии, фаза ровно от −180°, шаг 10° (15/20/30°) и наименьший подходящий шаг дБ
                (1/2/2,5/4/5/10 — у образца 3 2,5);
      "split" — нули на разной высоте: обе шкалы с нижней линии (L_min и −180°), шаги 1|2 дБ и 5|10°, деления
                обеих шкал на линиях сетки."""
    import numpy as np
    import tauplot as tp
    from matplotlib.lines import Line2D

    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    v, rows = d["variant"], d["rows"]
    K, a2, a1 = v["K"], v["a2"], v["a1"]
    W = lambda w: K / (1 - a2 * w ** 2 + 1j * a1 * w)
    we = np.array([r["w"] for r in rows])
    Ae = np.array([r["A"] for r in rows])
    pe = np.array([r["phi"] for r in rows])
    Le = np.array([r["L"] for r in rows])
    # рисунок на весь альбомный А4: поля 5 мм (меньше — принтер обрежет), подпись рисунка ≈ 10 мм
    PAGE_W, PAGE_H = 287, 190
    res = {}

    def legend(sh, items, loc):
        hs, ls = [], []
        for kind, color, text in items:
            hs.append(Line2D([], [], color=color, lw=1.8) if kind == "line" else
                      Line2D([], [], color=color, marker=kind, ls="none", ms=5))
            ls.append(text)
        lg = sh.ax.legend(hs, ls, loc=loc, fontsize=11, framealpha=1, edgecolor="black", fancybox=False)
        lg.set_zorder(8)

    def check(sh, key):
        need = sh.need
        if need["right"] > sh.pw + sh.margins[2] or need["top"] > sh.ph + sh.margins[3]:
            print(f"ВНИМАНИЕ: {key}: стрелка/подпись оси не влезает в поля", need, sh.margins)

    # ---------- АЧХ
    # Рисунок строго PAGE_W × PAGE_H: сетка = всё, что осталось за вычетом узких полей под числа и короткие стрелки;
    # клетка = место / число клеток (может быть дробной), подписи осей — у концов коротких стрелок.
    M = (9, 12, 8, 8)  # снизу — ряд чисел ω и под ним подпись «ω, рад/с»
    wmax = we.max()
    ax_x = tp.plan_exact(0, wmax, PAGE_W - M[0] - M[2])  # до последней точки, без запаса
    wc = np.linspace(0, ax_x.hi, 4000)
    Ac = np.abs(W(wc))
    ax_y = tp.plan_exact(0, max(Ac.max(), Ae.max()) * 1.005, PAGE_H - M[1] - M[3])  # до пика кривой
    sh = tp.Sheet(ax_x, ax_y, M)
    sh.ax.plot(wc, Ac, color=tp.CURVE, lw=1.8, zorder=4)
    sh.ax.plot(we, Ae, "o", color=tp.POINT, ms=5, zorder=6)
    sh.crossed_axes("ω, рад/с", "A(ω)", arrow=6, xarrow=7, obstacles=[sh.curve_mm(wc, Ac)],
                    markers=np.column_stack(sh.to_mm(we, Ae)), all_numbers=True)
    legend(sh, [("line", tp.CURVE, "A(ω)"), ("o", tp.POINT, "A(ωᵢ)")], "upper right")
    check(sh, "ach")
    res["ach"] = sh.save(outdir / "ach.png")

    # ---------- ФЧХ (ось ω — сверху, на уровне φ = 0)
    M = (11, 3, 8, 12)
    ax_x = tp.plan_exact(0, wmax, PAGE_W - M[0] - M[2])
    ax_y = tp.plan_exact(-180, 0, PAGE_H - M[1] - M[3])  # ровно −180…0°
    sh = tp.Sheet(ax_x, ax_y, M)
    phc = np.degrees(np.unwrap(np.angle(W(wc))))
    sh.ax.plot(wc, phc, color=tp.CURVE, lw=1.8, zorder=4)
    sh.ax.plot(we, pe, "o", color=tp.POINT, ms=5, zorder=6)
    sh.crossed_axes("ω, рад/с", "φ(ω), град", xnum="above", arrow=10, xarrow=7,
                    obstacles=[sh.curve_mm(wc, phc)], markers=np.column_stack(sh.to_mm(we, pe)), all_numbers=True)
    legend(sh, [("line", tp.CURVE, "φ(ω)"), ("o", tp.POINT, "φ(ωᵢ)")], "center right")
    check(sh, "fch")
    res["fch"] = sh.save(outdir / "fch.png")

    # ---------- АФЧХ: квадратная клетка, один шаг по Re и Im, разметка частот, без легенды
    # не на весь лист (просьба пользователя): 1–2 свободные клетки вокруг кривой, масштаб — максимальный
    M = (4, 3, 8, 8)
    wf = np.concatenate([[0.0], np.logspace(-4, 3, 20000)])
    Wf = W(wf)
    Ue, Ve = Ae * np.cos(np.radians(pe)), Ae * np.sin(np.radians(pe))
    xlo, xhi = min(Wf.real.min(), Ue.min()), max(Wf.real.max(), Ue.max())
    ylo = min(Wf.imag.min(), Ve.min())
    SUB = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")
    labels = [f"ω{str(i).translate(SUB)} = {w:.3f}".replace(".", ",") for i, w in enumerate(we)] + ["ω = 0"]
    # если подписи частот не помещаются (точки слиплись у начала координат) — больше запаса клеток вокруг кривой
    tries = ((1, 1, 1, 2), (2, 1, 2, 3), (3, 1, 3, 5), (4, 2, 4, 7), (5, 2, 5, 9))
    for mc in tries:  # сначала минимальный запас клеток вокруг кривой, больше — только если подписи не влезли
        ax_x, ax_y, _ = tp.plan_square_exact(xlo, xhi, ylo, 0, PAGE_W - M[0] - M[2], PAGE_H - M[1] - M[3],
                                             margin_cells=mc, fill=False)
        sh = tp.Sheet(ax_x, ax_y, M)
        sh.ax.plot(Wf.real, Wf.imag, color=tp.CURVE, lw=1.8, zorder=4)
        sh.ax.plot(Ue, Ve, "o", color=tp.POINT, ms=5, zorder=6)
        curve = sh.curve_mm(Wf.real, Wf.imag)
        pts = np.column_stack(sh.to_mm(Ue, Ve))
        Xk, Yk = sh.to_mm(K, 0)
        pts_all = np.vstack([pts, [[float(Xk), float(Yk)]]])
        sh.reserve_numbers(xnum="above")  # все деления осей остаются, подписи частот их обходят
        # «гроздь» точек у начала координат (ω → ∞) — столбиком внутри петли, остальные — обычной раскладкой
        O = np.array([float(v) for v in sh.to_mm(0, 0)])
        near = [i for i in range(len(pts)) if np.hypot(*(pts[i] - O)) < 8]
        leaders = sh.stack_labels(pts[near], [labels[i] for i in near], O, curves=curve) if len(near) >= 3 else None
        if leaders is None:
            near, leaders = [], []
        rest = [i for i in range(len(labels)) if i not in near]
        placed = sh.place_labels(pts_all[rest], [labels[i] for i in rest], [curve] + leaders, markers_mm=pts, size=9)
        unplaced = [labels[rest[k]] for k, r in enumerate(placed) if r is None]
        if not unplaced or mc == tries[-1]:
            break
        import matplotlib.pyplot as plt
        plt.close(sh.fig)
    sh.crossed_axes("Re W(jω)", "Im W(jω)", xnum="above", obstacles=[curve], markers=pts, arrow=6, xarrow=7,
                    all_numbers=True)
    check(sh, "afch")
    res["afch"] = sh.save(outdir / "afch.png")
    res["afch"]["unplaced"] = unplaced

    # ---------- ЛАЧХ + ЛФЧХ: общая ось ω, нули левой и правой шкал на одной линии
    # ось ω обрезана по точкам: от линии сетки перед первой точкой до линии после последней (0,03 … 0,1 … 1 … 6)
    M = (12, 7, 13, 8)
    wlo, whi = tp.log_nice(we.min()), tp.log_nice(we.max(), up=True)
    k0, k1 = math.log10(wlo), math.log10(whi)
    wl = np.logspace(k0, k1, 4000)
    Wl = W(wl)
    Lc = 20 * np.log10(np.abs(Wl))
    phl = np.degrees(np.unwrap(np.angle(Wl)))
    Lmax = max(Lc.max(), Le.max()) + 0.3  # верх шкалы — до пика ЛАЧХ, без пустых клеток
    Lneed = min(Le.min(), float(Lc.min())) - 0.3  # кривая до правого края не уходит под сетку
    AH = PAGE_H - M[1] - M[3]
    # Шкалы (все деления обеих шкал — на линиях сетки, подписана каждая линия):
    #  nice  — шаг на клетку 1|2 дБ и 5|10° (просьба пользователя 03.10.2026), нули на одной линии; z клеток ниже
    #          нуля: z·шаг_L ≥ |L_min| и z·шаг_φ ≥ 180° (фаза может уйти ниже −180°); из пар — с наименьшей долей
    #          пустых клеток (уход фазы ниже −180° — вдвое дороже), клетка ≥ 4,5 мм; при равенстве — мельче шаг;
    #  exact — нули на одной линии, фаза ровно от −180° (z·шаг_φ = 180°, шаг_φ 10/15/20/30°), шаг L — наименьший
    #          из 1/2/2,5/4/5/10 дБ, при котором точки помещаются (у образца 3 2,5 дБ / 10°, −45…+7,5 дБ);
    #  split — нули на разной высоте: обе шкалы начинаются с нижней линии (L_min по шагу и −180°), шаг 1|2 дБ и
    #          5|10°, клеток — сколько нужно большей из шкал (меньше пустых клеток), фаза выше 0° — сколько осталось.
    cands = []
    if lach_mode == "exact":
        for sp in (10, 15, 20, 30):
            z = int(180 / sp)
            for sl in (1, 2, 2.5, 4, 5, 10):
                if z * sl < -Lneed:
                    continue
                m = max(1, math.ceil(Lmax / sl - 1e-9))
                if AH / (z + m) >= 4.5:
                    cands.append((0, sl, sp, z, m, z + m, AH / (z + m)))
                    break
            if cands:
                break
    else:
        for sl in (1, 2):
            for sp in (5, 10):
                zL = math.ceil(-Lneed / sl - 1e-9)
                m = max(1, math.ceil(Lmax / sl - 1e-9))
                if lach_mode == "split":
                    n = max(zL + m, math.ceil(180 / sp - 1e-9))
                    z = zL
                    m = n - z
                    empty = (n - (Lmax - Lneed) / sl) / n + (n - 180 / sp) / n
                else:
                    z = max(math.ceil(180 / sp - 1e-9), zL)
                    n = z + m
                    # пустые клетки шкалы фазы ниже −180° смотрятся хуже пустого низа шкалы дБ — вдвое дороже
                    empty = (z - (-Lneed) / sl) / z + 2 * (z - 180 / sp) / z
                cell = AH / n
                if cell < 4.5:
                    continue
                cands.append((round(empty, 6), sl, sp, z, m, n, cell))
    if not cands:  # очень глубокая ЛАЧХ: крупный шаг, лишь бы влезло
        sl, sp = 5, 10
        z = max(18, math.ceil(-Lneed / sl)); m = max(1, math.ceil(Lmax / sl)); n = z + m; cell = AH / n
    else:
        _, sl, sp, z, m, n, cell = min(cands)
    P = sp
    ylabel_k = 1
    dec = k1 - k0                             # число декад (дробное при обрезке)
    ax_y = tp.Axis(-z * sl, m * sl, sl, n, cell)
    logx = dict(lo=wlo, hi=whi, mm=PAGE_W - M[0] - M[2])  # сетка на всю ширину листа
    sh = tp.Sheet(tp.Axis(0, 1, 1, 1, 1), ax_y, M, logx=logx, ylabel_k=ylabel_k)
    sh.ax.plot(wl, Lc, color=tp.CURVE2, lw=1.8, zorder=4)
    sh.ax.plot(we, Le, "o", color=tp.POINT2, ms=4.5, zorder=6)
    ax2 = sh.ax.twinx()
    if lach_mode == "split":
        ph_lo, ph_hi = -180, -180 + n * sp    # фаза с нижней линии, её ноль — на своей линии сетки
    else:
        ph_lo, ph_hi = -z * sp, m * sp
    ax2.set_ylim(ph_lo, ph_hi)
    ax2.set_yticks([])
    for spn in ax2.spines.values():
        spn.set_visible(False)
    ax2.plot(wl, phl, color=tp.CURVE, lw=1.8, zorder=4)
    ax2.plot(we, pe, "s", color=tp.POINT, ms=4, zorder=6)
    phi_mm = lambda p: (np.asarray(p, float) - ph_lo) / (ph_hi - ph_lo) * sh.ph
    cL = sh.curve_mm(wl, Lc)
    cP = np.column_stack([sh.to_mm(wl, 0)[0], phi_mm(phl)])
    cP = cP[(cP[:, 1] >= 0) & (cP[:, 1] <= sh.ph)]
    marks = np.vstack([np.column_stack(sh.to_mm(we, Le)), np.column_stack([sh.to_mm(we, 0)[0], phi_mm(pe)])])
    wa = 10.0 ** (k0 + 0.35 * dec)
    xa = float(sh.to_mm(wa, 0)[0])
    aL = [xa, float(sh.to_mm(wa, 20 * math.log10(abs(W(wa))))[1])]
    aP = [xa, float(phi_mm(math.degrees(np.angle(W(wa)))))]
    sh.place_labels([aL, aP], ["L(ω)", "φ(ω)"], [cL, cP], markers_mm=marks, size=12, prefer=lambda i: (0, 1))
    rticks = [ph_lo + k * sp for k in range(n + 1)]      # деления фазы — на каждой линии сетки
    # подписаны все деления; число фазы на линии L = 0 (при общем нуле — 0°) стоит над стрелкой оси ω
    right = dict(label="φ(ω), град", ticks=rticks, dec=0, skip=[], to_mm=lambda val: phi_mm(val))
    sh.crossed_axes("ω, рад/с", "L(ω), дБ", x_at=0, xnum="bottom", obstacles=[cL, cP], markers=marks,
                    right=right, arrow=6, xarrow=6, all_numbers=True)
    legend(sh, [("line", tp.CURVE2, "L(ω)"), ("o", tp.POINT2, "L(ωᵢ)"),
                ("line", tp.CURVE, "φ(ω)"), ("s", tp.POINT, "φ(ωᵢ)")], "lower left")
    check(sh, "lach")
    res["lach"] = sh.save(outdir / "lach.png")
    res["lach"]["scales"] = dict(mode=lach_mode, L_step=sl, L_label=ylabel_k * sl, phi_per_cell=round(sp, 4),
                                 phi_label=P, below_zero=z, above_zero=m, phi_range=[ph_lo, ph_hi])
    return res


def cmd_check(path):
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    worst = 0
    for r in d["rows"]:
        eA = abs(r["A"] / r["A_theory"] - 1) * 100
        eP = abs(r["phi"] - r["phi_theory"])
        worst = max(worst, eA)
        print(f"T={r['T']:8.3f} w={r['w']:6.3f}  A={r['A']:7.3f} (теор {r['A_theory']:7.3f}, {eA:4.1f}%)  "
              f"φ={r['phi']:8.2f} (теор {r['phi_theory']:8.2f}, Δ{eP:4.1f}°)")
    print(f"макс. отклонение A: {worst:.2f}%")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("make")
    m.add_argument("--name", required=True)
    m.add_argument("--group", default="")
    for k in ("a2", "a1", "K", "Um"):
        m.add_argument("--" + k, type=float)
    m.add_argument("--seed", type=int)
    m.add_argument("--dec", type=int, default=3,
                   help="знаков после запятой у всех величин (по умолчанию 3; 0 — прежнее округление «с экрана»)")
    m.add_argument("--out", required=True)
    m.add_argument("--teacher")
    m.add_argument("--title-page", action="store_true")
    m.add_argument("--no-mathcad", action="store_true")
    m.add_argument("--lach-variants", action="store_true",
                   help="ещё 2 комплекта графиков: ЛАФЧХ с разными нулями / с общим нулём и дробным шагом фазы")
    b = sub.add_parser("batch")
    b.add_argument("csv")
    b.add_argument("--out", required=True)
    b.add_argument("--teacher")
    b.add_argument("--title-page", action="store_true")
    b.add_argument("--no-mathcad", action="store_true")
    b.add_argument("--lach-variants", action="store_true")
    c = sub.add_parser("check")
    c.add_argument("json")
    a = ap.parse_args()

    if a.cmd == "check":
        return cmd_check(a.json)
    people = []
    if a.cmd == "make":
        people.append(dict(name=a.name, group=a.group, a2=a.a2, a1=a.a1, K=a.K, Um=a.Um, seed=a.seed, dec=a.dec))
    else:
        txt = Path(a.csv).read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(txt.splitlines()[0], delimiters=";,\t")
        for row in csv.DictReader(txt.splitlines(), dialect=dialect):
            row = {k.strip(): (v or "").strip() for k, v in row.items() if k}
            if not row.get("name"):
                continue
            num = lambda k: float(row[k].replace(",", ".")) if row.get(k) else None
            people.append(dict(name=row["name"], group=row.get("group", ""), a2=num("a2"), a1=num("a1"),
                               K=num("K"), Um=num("Um"), seed=row.get("seed") or None,
                               dec=row.get("dec") or 3))
    summary = []
    for pd in people:
        d, folder, rep = cmd_make_one(pd, a.out, not a.no_mathcad, a.teacher, a.title_page,
                                      getattr(a, "lach_variants", False))
        v = d["variant"]
        line = (f"{d['name']}: a2={v['a2']} a1={v['a1']} K={v['K']} Um={v['Um']} | T0={d['rows'][0]['T']} "
                f"q={d['style']['q']} n={len(d['rows'])} | {folder}")
        if rep:
            line += f" | Mathcad ошибок: {rep.get('errors')}, графиков: {len(rep.get('plots', {}))}"
        print(line)
        summary.append(line)


if __name__ == "__main__":
    main()
