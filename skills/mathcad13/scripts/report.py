#!/usr/bin/env python3
"""report.py - Word (.docx) report from a Mathcad 13 worksheet.

Text regions become real Word paragraphs (headings keep their level), formulas and plots become
pictures rendered by Mathcad itself. Mathcad renders at screen resolution (96 dpi), so the worksheet
is first copied with math fonts enlarged by --scale and the pictures are shrunk back in Word.

  report.py IN.xmcd OUT.docx [--scale 2.5] [--math-zoom 1.4] [--font "Times New Roman"] [--size 14]
                             [--title "..."] [--keep DIR] [--no-recalc]

  --math-zoom  size of formulas relative to Mathcad's 10 pt math font (1.4 ~ matches 14 pt text)
  --keep DIR   keep the intermediate hi-res worksheet and HTML export in DIR

Page: A4, margins 30/15/20/20 mm (left/right/top/bottom, GOST 7.32 style).
Requires python-docx and Mathcad 13 (driven through mc.ps1).
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import xmcd  # noqa: E402

MC_PS1 = os.path.join(HERE, "mc.ps1")
HEADING = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3, "Title": 0}


def run_mc(path, save_as, no_recalc=False):
    cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", MC_PS1, "-Path", path, "-SaveAs", save_as]
    if no_recalc:
        cmd.append("-NoRecalc")
    out = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    try:
        res = json.loads(out.stdout[out.stdout.find("{"):])
    except ValueError:
        raise RuntimeError("mc.ps1 failed:\n" + out.stdout + out.stderr)
    if not res.get("ok"):
        raise RuntimeError("mc.ps1: " + str(res.get("error")))
    return res


def hires_copy(src, dst, scale):
    text = xmcd.strip_digest(open(src, encoding="utf-8").read())
    text = re.sub(r'(<mathStyle [^>]*font-size=")(\d+(?:\.\d+)?)(")',
                  lambda m: m.group(1) + f"{float(m.group(2)) * scale:g}" + m.group(3), text)
    with open(dst, "w", encoding="utf-8") as fh:
        fh.write(text)


def html_images(html_path):
    """Return, per region div in document order, the image path or None."""
    html = open(html_path, encoding="utf-8", errors="replace").read()
    base = os.path.dirname(html_path)
    out = []
    for div in re.finditer(r"<div\s[^>]*>(.*?)</div>", html, re.S):
        m = re.search(r'<img src="([^"]+)"', div.group(1))
        out.append(os.path.normpath(os.path.join(base, m.group(1))) if m else None)
    return out


def worksheet_regions(path):
    root = ET.parse(path).getroot()
    regions = next(k for k in root if xmcd.local(k.tag) == "regions")
    regs = []
    for r in regions:  # top level only: the HTML export has one div per top-level region
        if xmcd.local(r.tag) != "region":
            continue
        kids = [k for k in r if xmcd.local(k.tag) != "rendering"]
        kind = xmcd.local(kids[0].tag) if kids else "?"
        paras = []
        if kind == "text":
            for p in kids[0].iter():
                if xmcd.local(p.tag) == "p":
                    text = re.sub(r"\s+", " ", "".join(p.itertext())).strip()
                    paras.append((p.get("style", "Normal"), text, p_runs(p)))
        regs.append(dict(kind=kind, left=float(r.get("left", 0)), top=float(r.get("top", 0)),
                         width=float(r.get("width", 0)), height=float(r.get("height", 0)), paras=paras))
    return regs


INLINE = {"b": "bold", "i": "italic", "u": "underline", "sub": "sub", "sup": "sup"}


def p_runs(el, fmt=frozenset()):
    """Mathcad text paragraph -> [(text, {bold, italic, underline, sub, sup})] keeping inline markup."""
    tag = xmcd.local(el.tag)
    inner = fmt | {INLINE[tag]} if tag in INLINE else fmt
    out = [("\n", fmt)] if tag == "br" else [("\t", fmt)] if tag == "tab" else []
    if el.text:
        out.append((el.text, inner))
    for ch in el:
        out += p_runs(ch, inner)
        if ch.tail:
            out.append((ch.tail, inner))
    return out


def add_runs(p, runs, italic=False):
    for i, (text, fmt) in enumerate(runs):
        text = re.sub(r"[ \t\r\n]+", " ", text) if text not in ("\n", "\t") else text
        if i == 0:
            text = text.lstrip()
        if not text:
            continue
        r = p.add_run(text)
        r.bold = "bold" in fmt or None
        r.italic = ("italic" in fmt or italic) or None
        r.underline = "underline" in fmt or None
        if "sub" in fmt:          # both map to w:vertAlign - assigning None to the other one would clear it
            r.font.subscript = True
        elif "sup" in fmt:
            r.font.superscript = True


def set_fonts(doc, font, size):
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = font
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), font)
        rfonts.set(qn("w:cs"), font)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):  # theme fonts win otherwise
            rfonts.attrib.pop(qn(attr), None)
        st.font.color.rgb = None
        if name == "Normal":
            st.font.size = Pt(size)
        else:
            st.font.size = Pt(size + {"Heading 1": 2, "Heading 2": 1, "Heading 3": 0, "Title": 4}[name])
            st.font.bold = True
    pf = doc.styles["Normal"].paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)


def build_docx(regs, images, plot_images, out, a):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(30), Mm(15), Mm(20), Mm(20)
    set_fonts(doc, a.font, a.size)
    max_w = (210 - 30 - 15) / 25.4
    if a.title:
        p = doc.add_paragraph(a.title, style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    comments = attach_comments(regs)
    is_comment = {c for cs in comments.values() for c in cs}
    order = sorted(range(len(regs)), key=lambda i: (regs[i]["top"], regs[i]["left"]))
    for i in order:
        if i in is_comment:
            continue
        r, img = regs[i], images[i] if i < len(images) else None
        if r["kind"] == "plot":  # plot labels use the math fonts, so take plots from the unscaled export
            img = plot_images[i] if i < len(plot_images) else img
        if r["kind"] == "pageBreak":
            doc.add_page_break()
            continue
        if r["kind"] == "text":
            for style, t, runs in r["paras"]:
                if not t:
                    continue
                lvl = HEADING.get(style)
                p = doc.add_paragraph(style="Title" if lvl == 0 else (f"Heading {lvl}" if lvl else "Normal"))
                add_runs(p, runs)
                if re.match(r"(Рисунок|Рис\.)\s", t):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if not img or not os.path.exists(img):
            continue
        from PIL import Image
        with Image.open(img) as im:
            wpx = im.size[0]
        if r["kind"] == "plot":
            width = min(wpx / 96.0 * a.plot_zoom, max_w)
        else:
            width = min(wpx / 96.0 / a.scale * a.math_zoom, max_w)
        p = doc.add_paragraph()
        if r["kind"] == "plot":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(width))
        for c in comments.get(i, []):                          # comment written to the right of it
            if any(t for _, t, _ in regs[c]["paras"]):
                p.paragraph_format.tab_stops.add_tab_stop(Mm(a.comment_tab))
                p.add_run("\t")
                for _, _, runs in regs[c]["paras"]:
                    add_runs(p, runs, italic=True)
    doc.save(out)


def attach_comments(regs):
    """Text regions standing to the right of a formula (vertical overlap) are that formula's comment."""
    comments = {}
    for i, r in enumerate(regs):
        if r["kind"] != "text":
            continue
        best, best_ov = None, 0.0
        for j, m in enumerate(regs):
            if m["kind"] == "text" or m["left"] + m["width"] > r["left"] + 4:
                continue
            ov = min(r["top"] + max(r["height"], 12), m["top"] + m["height"]) - max(r["top"], m["top"])
            if ov > best_ov:
                best, best_ov = j, ov
        if best is not None:
            comments.setdefault(best, []).append(i)
    return comments


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inp")
    ap.add_argument("out")
    ap.add_argument("--scale", type=float, default=2.5)
    ap.add_argument("--math-zoom", type=float, default=1.4)
    ap.add_argument("--plot-zoom", type=float, default=1.0)
    ap.add_argument("--font", default="Times New Roman")
    ap.add_argument("--size", type=float, default=14)
    ap.add_argument("--title", default="")
    ap.add_argument("--comment-tab", type=float, default=75, help="tab position (mm) of side comments")
    ap.add_argument("--keep", default="")
    ap.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args()
    work = a.keep or tempfile.mkdtemp(prefix="mcreport_")
    os.makedirs(work, exist_ok=True)
    src = os.path.abspath(a.inp)
    # next to the original, so relative data-file paths (READPRN("data.prn") etc.) still resolve
    hi = os.path.join(os.path.dirname(src), "~report_" + os.path.basename(src))
    try:
        hires_copy(src, hi, a.scale)
        html = os.path.join(work, "export.htm")
        res = run_mc(hi, html, a.no_recalc)
        regs = worksheet_regions(src)
        imgs = html_images(html)
        plot_imgs = []
        if any(r["kind"] == "plot" for r in regs):
            html_lo = os.path.join(work, "plots", "export.htm")
            run_mc(src, html_lo, a.no_recalc)
            plot_imgs = html_images(html_lo)
        if len(imgs) != len(regs):
            print(f"warning: {len(imgs)} exported blocks vs {len(regs)} regions - mapping may be off", file=sys.stderr)
        build_docx(regs, imgs, plot_imgs, a.out, a)
        print(f"wrote {a.out}: {len(regs)} regions, {res.get('error_count', 0)} region(s) with errors")
    finally:
        if a.keep:
            shutil.copy(hi, os.path.join(work, "hires.xmcd"))
        if os.path.exists(hi):
            os.remove(hi)
        if not a.keep:
            shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
