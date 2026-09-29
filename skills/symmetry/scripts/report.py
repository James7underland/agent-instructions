#!/usr/bin/env python3
"""report.py — отчёт по кейсу Symmetry: Word (.docx) со схемой и таблицей потоков, и/или Excel.

  python report.py CASE.vsym OUT.docx [--units SI] [--pfd PFD.png | --auto-pfd] [--title "..."]
                   [--props VapFrac,T,P,MoleFlow,MassFlow] [--comp mole|mass|none] [--streams Feed,Gas,Liq]
                   [--xlsx OUT.xlsx] [--extra "/H1.InQ.Energy@kW=Тепловая нагрузка H1"]...
Таблица «материального баланса» — потоки по столбцам, свойства и состав по строкам (как принято в курсовых).
--auto-pfd: схема рисуется pfd.py (топология из кейса). --extra: отдельные величины (Q, N, КПД) списком после таблицы.
Шрифт Times New Roman 12 (таблица 10), поля 2/1/2/3 см (ГОСТ 7.32 для учебных работ). Числа — через запятую.
"""
from __future__ import annotations

import argparse
import os
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sym  # noqa: E402

NAMES = {"VapFrac": "Доля пара", "T": "Температура", "P": "Давление", "MoleFlow": "Мольный расход",
         "MassFlow": "Массовый расход", "VolumeFlow": "Объёмный расход", "StdLiqVolumeFlow": "Объём. расход (ст. жидк.)",
         "StdGasVolumeFlow": "Объём. расход газа (ст. усл.)", "MolecularWeight": "Молярная масса",
         "MassDensity": "Плотность", "H": "Энтальпия (мольн.)", "Energy": "Поток энтальпии", "Cp": "Теплоёмкость",
         "Viscosity": "Вязкость", "ZFactor": "Коэф. сжимаемости Z"}


def fmt(v):
    if v is None:
        return "—"
    if isinstance(v, float):
        a = abs(v)
        s = f"{v:.4f}" if a < 1 else (f"{v:.3f}" if a < 100 else (f"{v:.1f}" if a < 1e5 else f"{v:.4g}"))
        return s.replace(".", ",")
    return str(v)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("case"); ap.add_argument("out")
    ap.add_argument("--units", default="SI"); ap.add_argument("--pfd"); ap.add_argument("--auto-pfd", action="store_true")
    ap.add_argument("--title", default=""); ap.add_argument("--props", default="VapFrac,T,P,MoleFlow,MassFlow,MolecularWeight,MassDensity")
    ap.add_argument("--comp", default="mole", choices=["mole", "mass", "none"]); ap.add_argument("--streams")
    ap.add_argument("--xlsx"); ap.add_argument("--extra", action="append", default=[])
    a = ap.parse_args([sym.unmangle(x) for x in sys.argv[1:]])

    pfd = a.pfd
    if a.auto_pfd and not pfd:
        pfd = os.path.join(tempfile.gettempdir(), "sym_report_pfd.png")
        subprocess.run([sys.executable, os.path.join(os.path.dirname(__file__), "pfd.py"), a.case, pfd, "--values",
                        "--units", a.units], check=True, env=dict(os.environ, PYTHONIOENCODING="utf-8"))

    eng = sym.Engine(echo=lambda *x: None)
    sym.open_case(eng, a.case)
    eng.eval(f"units {a.units}"); eng.msgs.clear()
    props = a.props.split(",")
    rows, units, cmps = sym.stream_table(eng, props, a.comp)
    if a.streams:
        want = ["/" + s.lstrip("/") for s in a.streams.split(",")]
        rows = sorted([r for r in rows if r["Stream"] in want], key=lambda r: want.index(r["Stream"]))
    extras = []
    for e in a.extra:
        spec, _, label = e.partition("=")
        p, u = sym.parse_path_unit(sym.unmangle(spec))
        try:
            v, uu = eng.value(p, u)
        except Exception:
            v, uu = None, u or ""
        extras.append((label or p, v, uu))
    if a.xlsx:
        sym.write_table(rows, units, a.xlsx)

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1), Cm(2), Cm(2)
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
    if a.title:
        h = doc.add_paragraph(a.title); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.runs[0].bold = True
    n_fig = 0
    if pfd and os.path.exists(pfd):
        doc.add_picture(pfd, width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        n_fig += 1
        c = doc.add_paragraph(f"Рисунок {n_fig} – Технологическая схема (расчёт Symmetry)")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Таблица 1 – Параметры материальных потоков (набор единиц {a.units})")
    names = [r["Stream"].lstrip("/") for r in rows]
    comp_keys = [k for k in (rows[0].keys() if rows else []) if k.startswith(("x_", "w_"))]
    body = [(f"{NAMES.get(p, p)}" + (f", {units[p]}" if units.get(p) else ""), [r.get(p) for r in rows]) for p in props]
    if comp_keys:
        body.append(("Состав, " + ("мол. доли" if a.comp == "mole" else "масс. доли"), None))
        for k in comp_keys:
            body.append((f"  {k[2:]}", [r.get(k) for r in rows]))
    # делим широкие таблицы по 6 потоков
    chunk = 6
    for i0 in range(0, max(1, len(names)), chunk):
        part = names[i0:i0 + chunk]
        t = doc.add_table(rows=1 + len(body), cols=1 + len(part))
        t.style = "Table Grid"
        t.cell(0, 0).text = "Параметр"
        for j, nm in enumerate(part):
            t.cell(0, j + 1).text = nm
        for i, (lab, vals) in enumerate(body, start=1):
            t.cell(i, 0).text = lab
            if vals is None:
                continue
            for j in range(len(part)):
                t.cell(i, j + 1).text = fmt(vals[i0 + j])
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10)
        doc.add_paragraph("")
    if extras:
        doc.add_paragraph("Расчётные величины:")
        for lab, v, u in extras:
            doc.add_paragraph(f"{lab} = {fmt(v)} {u}".rstrip(), style="List Bullet")
    doc.save(a.out)
    print(f"docx -> {a.out} (потоков {len(rows)}, компонентов {len(cmps)})")
    sym.hard_exit(0)


if __name__ == "__main__":
    main()
