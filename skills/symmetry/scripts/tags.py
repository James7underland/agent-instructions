#!/usr/bin/env python3
"""tags.py — перечень сигналов (тегов) между моделью Symmetry и ПЛК: сбор по модели, адреса, экспорт.

  scan CASE.vsym -o tags.csv [--mb-r 0 --mb-w 1000] [--streams]   теги по объектам модели (см. ниже)
  from-xml OPC_DA.XML -o tags.csv                                  OPC-конфигурация Symmetry → карта тегов
  to-xml tags.csv -o OPC.xml [--group Main]                        карта → XML для OPC DA сервера Symmetry
  mb tags.csv [-o tags.csv] [--mb-r 0 --mb-w 1000]                 раздать адреса Modbus (float = 2 рег.)
  st tags.csv -o gvl.st [--name GVL_SYM] [--ua-prefix ...]         объявления переменных ПЛК (МЭК 61131-3 ST)
  xlsx tags.csv -o io_list.xlsx                                    перечень входов/выходов для программиста ПЛК
  browse opc.tcp://host:4840 [--depth 4] [--root "ns=2;s=..."]     узлы OPC UA сервера (NodeId для колонки ua)

scan: Controller Indicator/Off → <имя>.PV (R); Controller Auto/Cascade → .PV R, .SP W, .OP R;
Controller Manual с рампой выхода (позиционер) → <клапан>.CMD W (OPTarget); Manual без рампы → .OP W;
Valve → <имя>.POS R (Actual_Pos), .CMD W (%Opening), если открытие не подключено к регулятору;
Pump → .RUN W (Switch On/Off), .RUNFB R, .SPEED R; Flash/Tank → .LEVEL R (Liq0Level%), .P R;
--streams: граничные потоки (без источника/приёмника) → .P, .T, .F (массовый расход) R.
Колонки карты: tag;path;unit;dir;type;mb;expr;ua;desc (формат plcbridge.py).
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import xml.etree.ElementTree as ET

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
COLS = ["tag", "path", "unit", "dir", "type", "mb", "expr", "ua", "desc"]


# ---------------------------------------------------------------- чтение/запись карты
def read_map(path):
    raw = open(path, encoding="utf-8-sig").read().splitlines()
    delim = ";" if raw[0].count(";") >= raw[0].count(",") else ","
    rows = []
    for r in csv.DictReader([l for l in raw if l.strip() and not l.lstrip().startswith("#")], delimiter=delim):
        r = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
        if r.get("tag"):
            r["dir"] = (r.get("dir") or "R").upper()[:1]
            r["type"] = (r.get("type") or "float").lower()
            rows.append({c: r.get(c, "") for c in COLS})
    return rows


def write_map(rows, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS, delimiter=";")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in COLS})
    print(f"# {len(rows)} тегов -> {path}")


def nregs(r):
    return 2 if r["type"] == "float" else 1


def assign_mb(rows, r0=0, w0=1000, keep=True):
    used = set()
    if keep:
        for r in rows:
            if str(r.get("mb", "")).strip():
                used |= set(range(int(r["mb"]), int(r["mb"]) + nregs(r)))
    nxt = {"R": r0, "W": w0}
    for r in rows:
        if keep and str(r.get("mb", "")).strip():
            continue
        a = nxt[r["dir"]]
        while any(x in used for x in range(a, a + nregs(r))):
            a += 1
        r["mb"] = str(a)
        used |= set(range(a, a + nregs(r)))
        nxt[r["dir"]] = a + nregs(r)
    return rows


# ---------------------------------------------------------------- scan
UNIT_RE = re.compile(r"=\s*(-?[\d.,eE+-]+|None|-12321)\s*\*?\s*(\S*)\s*$")


def _unit_of(txt):
    """Единица из ответа движка на запрос переменной (последняя строка 'Тип = 50,0 * %')."""
    for line in reversed(str(txt or "").splitlines()):
        m = UNIT_RE.search(line.strip())
        if m:
            return m.group(2)
    return ""


def _conn(txt):
    m = re.search(r"Conn\. to:\s*(.*)", str(txt or ""))
    c = (m.group(1).strip() if m else "")
    return "" if c in ("-", "") else c


def _param(eng, path):
    t = str(eng.eval(path) or "")
    eng.msgs.clear()
    m = re.search(r"=\s*(.*?)\s*$", t.splitlines()[0]) if t else None
    return m.group(1).strip() if m else ""


def scan(case, streams=False):
    import sym
    eng = sym.Engine()
    eng.e.AddDynamicsSupport2()
    eng.flush()
    sym.open_case(eng, case)
    ops = [(n, t) for n, t in eng.children("/") if t.startswith("sim.unitop.")]
    rows = []
    driven = {}                        # путь переменной → кто ею управляет (регулятор)

    def q(p):
        t = eng.eval(p)
        eng.msgs.clear()
        return t

    def add(tag, path, dir, unit="", type="float", expr="", desc=""):
        rows.append({"tag": tag, "path": path, "unit": unit, "dir": dir, "type": type, "mb": "", "expr": expr,
                     "ua": "", "desc": desc})

    ctrls = [n for n, t in ops if t.endswith("Controller.Controller")]
    for n in ctrls:
        mode = _param(eng, f"/{n}.Mode")
        tin, tout = q(f"/{n}.In"), q(f"/{n}.Out")
        pv_src, op_dst = _conn(tin), _conn(tout)
        u_pv = _unit_of(tin)
        if op_dst:
            driven[op_dst] = n
        ramp = _param(eng, f"/{n}.Setpoint_Ramp_Mode")
        if mode in ("Indicator", "Off") or (pv_src and not op_dst):
            if pv_src:
                add(f"{n}.PV", f"/{n}.In", "R", u_pv, desc=f"датчик: {pv_src}")
        elif mode in ("Automatic", "Cascade"):
            add(f"{n}.PV", f"/{n}.In", "R", u_pv, desc=f"PV регулятора: {pv_src}")
            add(f"{n}.SP", f"/{n}.Target", "W", u_pv, desc="задание регулятора Symmetry")
            add(f"{n}.OP", f"/{n}.OP", "R", "%", desc=f"выход → {op_dst}")
        elif mode == "Manual" and op_dst:
            dev = op_dst.split(".")[0].lstrip("/")
            if ramp == "Output":
                add(f"{dev}.CMD", f"/{n}.OPTarget", "W", "%", desc=f"команда позиционеру {n} → {op_dst}")
            else:
                add(f"{n}.OP", f"/{n}.OP", "W", "%", desc=f"ручной выход → {op_dst}")
    for n, t in ops:
        kind = t.rsplit(".", 1)[-1]
        if kind == "Valve":
            add(f"{n}.POS", f"/{n}.Actual_Pos", "R", "%", desc="фактическое положение клапана")
            if not any(k.startswith(f"/{n}.") for k in driven):
                add(f"{n}.CMD", f"/{n}.%Opening", "W", "%", desc="команда открытия (мгновенно; для времени хода — позиционер)")
        elif kind in ("PumpWithCurve",):
            add(f"{n}.RUN", f"/{n}.Switch", "W", "", "bool", "'On' if x else 'Off'", "пуск/стоп насоса")
            add(f"{n}.RUNFB", f"/{n}.Switch", "R", "", "bool", "", "насос в работе")
            add(f"{n}.SPEED", f"/{n}.PumpSpeed", "R", "rpm", desc="частота вращения")
        elif kind in ("SimpleFlash", "ThreePhaseSeparator", "Tank", "SimpleTank"):
            add(f"{n}.LEVEL", f"/{n}.Liq0Level%", "R", "%", desc="уровень (без КИП; лучше Controller-Indicator)")
            add(f"{n}.P", f"/{n}.Vap.P", "R", "kPa", desc="давление в аппарате")
    if streams:
        for n, t in ops:
            if not t.endswith("Stream_Material"):
                continue
            ti, to = q(f"/{n}.In"), q(f"/{n}.Out")
            if "Conn. to: -" in str(ti) or "Conn. to: -" in str(to):
                add(f"{n}.P", f"/{n}.In.P", "R", "kPa", desc="граничный поток: давление")
                add(f"{n}.T", f"/{n}.In.T", "R", "C", desc="граничный поток: температура")
                add(f"{n}.F", f"/{n}.In.MassFlow", "R", "kg/h", desc="граничный поток: массовый расход")
    # проверка: читается ли каждый путь
    bad = []
    for r in rows:
        try:
            v, _ = eng.value(r["path"], r["unit"] or None)
        except Exception:
            v = None
        if v is None:
            bad.append(r["tag"])
    if bad:
        print("# не читаются (проверьте):", ", ".join(bad))
    sym.hard_exit_after = True
    return rows


# ---------------------------------------------------------------- XML Symmetry OPC
def from_xml(path):
    rows = []
    for g in ET.parse(path).iter("Group"):
        for v in g.iter("Variable"):
            u = v.findtext("UnitName") or ""
            rows.append({"tag": v.findtext("OPCTag"), "path": v.findtext("SimPath"),
                         "unit": "" if u.lower() == "active set" else u,
                         "dir": "W" if (v.findtext("CanWrite") or "").lower() == "true" else "R",
                         "type": "float" if (v.findtext("VarType") or "Float").lower() in ("float", "double")
                         else ("bool" if (v.findtext("VarType") or "").lower() == "bool" else "int"),
                         "mb": "", "expr": "", "ua": "", "desc": f"группа {g.get('Name')}"})
    return rows


def to_xml(rows, out, group):
    main = ET.Element("Main")
    st = ET.SubElement(main, "Settings")
    ET.SubElement(st, "ActiveCaseId").text = "My case 3"
    vs = ET.SubElement(main, "Variables")
    gr = ET.SubElement(vs, "Group", Name=group)
    for r in rows:
        v = ET.SubElement(gr, "Variable")
        ET.SubElement(v, "OPCTag").text = r["tag"].replace("%", "%_")
        ET.SubElement(v, "SimPath").text = r["path"]
        ET.SubElement(v, "CanWrite").text = "True" if r["dir"] == "W" else "False"
        ET.SubElement(v, "VarType").text = "Float"
        ET.SubElement(v, "VarShape").text = "Scalar"
        ET.SubElement(v, "UnitName").text = r["unit"] or "Active Set"
    ET.indent(main)
    ET.ElementTree(main).write(out, encoding="utf-8", xml_declaration=True)
    print(f"# XML ({len(rows)} тегов, группа {group}) -> {out}")


# ---------------------------------------------------------------- ST / XLSX
def ident(tag):
    s = re.sub(r"[^A-Za-z0-9_]+", "_", tag.replace("%_", "").replace("%", "Pct")).strip("_")
    return ("_" + s) if s[0].isdigit() else s


def to_st(rows, out, name):
    L = [f"(* Переменные обмена с моделью Symmetry (plcbridge.py). Сгенерировано tags.py.",
         "   R — вход ПЛК (модель → ПЛК), W — выход ПЛК (ПЛК → модель).",
         "   Modbus: holding-регистры с 0, REAL = 2 регистра. *)",
         f"VAR_GLOBAL (* {name} *)"]
    for r in rows:
        ty = {"float": "REAL", "int": "INT", "bool": "BOOL"}[r["type"]]
        io = "вход" if r["dir"] == "R" else "выход"
        mb = f" MB {r['mb']}" if r.get("mb") else ""
        L.append(f"    {ident(r['tag']):28s}: {ty:5s}; (* {io}{mb}; {r['path']} [{r['unit']}] {r['desc']} *)")
    L.append("END_VAR")
    open(out, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(f"# ST -> {out}")


def to_xlsx(rows, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Перечень сигналов"
    head = ["№", "Тег", "Направление", "Тип ПЛК", "Modbus (holding)", "Регистров", "Переменная модели", "Ед.",
            "Пересчёт (x)", "OPC UA NodeId", "Описание"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 1):
        ws.append([i, r["tag"], "модель → ПЛК (вход)" if r["dir"] == "R" else "ПЛК → модель (выход)",
                   {"float": "REAL", "int": "INT", "bool": "BOOL"}[r["type"]],
                   int(r["mb"]) if r.get("mb") else None, nregs(r) if r.get("mb") else None, r["path"], r["unit"],
                   r["expr"], r["ua"] or f"ns=2;s={r['tag']}", r["desc"]])
    for col, w in zip("ABCDEFGHIJK", [5, 22, 22, 9, 10, 9, 30, 8, 18, 30, 50]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    print(f"# XLSX -> {out}")


# ---------------------------------------------------------------- browse
def browse(url, depth, root):
    from asyncua.sync import Client
    c = Client(url, timeout=5)
    c.connect()
    try:
        print("# namespaces:", c.get_namespace_array())
        start = c.get_node(root) if root else c.nodes.objects
        n = [0]

        def walk(node, d):
            for ch in node.get_children():
                n[0] += 1
                if n[0] > 3000:
                    return
                bn = ch.read_browse_name()
                nc = ch.read_node_class().name
                if bn.NamespaceIndex == 0 and bn.Name == "Server":
                    continue
                extra = ""
                if nc == "Variable":
                    try:
                        dv = ch.read_data_value()
                        extra = f" = {dv.Value.Value!r}"[:60] + f"  ({dv.Value.VariantType.name})"
                        al = ch.read_attribute(18).Value.Value      # UserAccessLevel
                        extra += "  RW" if al & 2 else "  R"
                    except Exception as ex:
                        extra = f"  <{type(ex).__name__}>"
                print("  " * d + f"{bn.Name}  [{nc}]  {ch.nodeid.to_string()}{extra}")
                if nc == "Object" and d < depth:
                    walk(ch, d + 1)
        walk(start, 0)
    finally:
        c.disconnect()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["scan", "from-xml", "to-xml", "mb", "st", "xlsx", "browse"])
    ap.add_argument("src")
    ap.add_argument("-o", "--out")
    ap.add_argument("--mb-r", type=int, default=0)
    ap.add_argument("--mb-w", type=int, default=1000)
    ap.add_argument("--streams", action="store_true")
    ap.add_argument("--group", default="Main")
    ap.add_argument("--name", default="GVL_SYM")
    ap.add_argument("--depth", type=int, default=4)
    ap.add_argument("--root")
    a = ap.parse_args([__import__("sym").unmangle(x) for x in sys.argv[1:]])
    if a.cmd == "browse":
        return browse(a.src, a.depth, a.root)
    if a.cmd == "scan":
        rows = assign_mb(scan(a.src, a.streams), a.mb_r, a.mb_w, keep=False)
        write_map(rows, a.out or "tags.csv")
        import sym
        sym.hard_exit(0)
    if a.cmd == "from-xml":
        return write_map(from_xml(a.src), a.out or "tags.csv")
    rows = read_map(a.src)
    if a.cmd == "to-xml":
        return to_xml(rows, a.out or "OPC.xml", a.group)
    if a.cmd == "mb":
        return write_map(assign_mb(rows, a.mb_r, a.mb_w), a.out or a.src)
    if a.cmd == "st":
        return to_st(rows, a.out or "gvl.st", a.name)
    if a.cmd == "xlsx":
        return to_xlsx(rows, a.out or "io_list.xlsx")


if __name__ == "__main__":
    main()
